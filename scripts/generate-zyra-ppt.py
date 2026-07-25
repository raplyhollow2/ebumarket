#!/usr/bin/env python3
"""Generate Zyra international product overview PPTX."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

FOREST = RGBColor(0x1C, 0x30, 0x24)
CREAM = RGBColor(0xF7, 0xF4, 0xEF)
WARM = RGBColor(0xEB, 0xE4, 0xD6)
MUTED = RGBColor(0x5C, 0x6B, 0x62)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x2F, 0x4A, 0x3A)
SOFT = RGBColor(0xD8, 0xD0, 0xC0)
CARD = RGBColor(0xFF, 0xFC, 0xF7)
TOTAL = 14
OUT = Path(__file__).resolve().parents[1] / "docs/presentations/Zyra-Product-Overview.pptx"


def set_run(run, size=18, bold=False, color=FOREST, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def fill_solid(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def rect(slide, l, t, w, h, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    fill_solid(s, color)
    return s


def round_rect(slide, l, t, w, h, color):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    fill_solid(s, color)
    try:
        s.adjustments[0] = 0.08
    except Exception:
        pass
    return s


def textbox(slide, l, t, w, h):
    return slide.shapes.add_textbox(l, t, w, h)


def set_first(tf, text, size=18, bold=False, color=FOREST, align=PP_ALIGN.LEFT, font="Calibri"):
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color, font=font)
    return tf


def add_para(tf, text, size=16, bold=False, color=FOREST, align=PP_ALIGN.LEFT, space_before=0, font="Calibri"):
    p = tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color, font=font)
    return p


def footer(slide, page, w):
    tb = textbox(slide, Inches(0.6), Inches(7.05), Inches(10), Inches(0.3))
    set_first(tb.text_frame, "Zyra  ·  Confidential", size=11, color=MUTED)
    tb2 = textbox(slide, Inches(11.5), Inches(7.05), Inches(1.3), Inches(0.3))
    set_first(tb2.text_frame, f"{page} / {TOTAL}", size=11, color=MUTED, align=PP_ALIGN.RIGHT)


def section_label(slide, text):
    tb = textbox(slide, Inches(0.7), Inches(0.45), Inches(8), Inches(0.35))
    set_first(tb.text_frame, text.upper(), size=12, bold=True, color=ACCENT)


def title(slide, text):
    tb = textbox(slide, Inches(0.7), Inches(0.75), Inches(11.5), Inches(0.8))
    set_first(tb.text_frame, text, size=32, bold=True, color=FOREST, font="Georgia")


def subtitle(slide, text):
    tb = textbox(slide, Inches(0.7), Inches(1.5), Inches(11), Inches(0.5))
    set_first(tb.text_frame, text, size=16, color=MUTED)


def bg_cream(slide, w, h):
    rect(slide, 0, 0, w, h, CREAM)
    rect(slide, 0, 0, Inches(0.18), h, FOREST)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    W, H = prs.slide_width, prs.slide_height

    # 1 Cover
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, FOREST)
    rect(s, 0, Inches(5.9), W, Inches(1.6), ACCENT)
    rect(s, Inches(0.7), Inches(2.55), Inches(1.4), Inches(0.08), WARM)
    tb = textbox(s, Inches(0.7), Inches(1.3), Inches(11), Inches(0.4))
    set_first(tb.text_frame, "PRODUCT OVERVIEW", size=14, bold=True, color=WARM)
    tb = textbox(s, Inches(0.7), Inches(1.8), Inches(11.5), Inches(1.0))
    set_first(tb.text_frame, "Zyra", size=60, bold=True, color=WHITE, font="Georgia")
    tb = textbox(s, Inches(0.7), Inches(2.85), Inches(11), Inches(0.9))
    tf = set_first(tb.text_frame, "Clothes that get a second life.", size=26, color=CREAM, font="Georgia")
    add_para(tf, "A trusted teen marketplace and donation platform for circular fashion.", size=16, color=SOFT, space_before=10)
    tb = textbox(s, Inches(0.7), Inches(6.2), Inches(8), Inches(0.8))
    tf = set_first(tb.text_frame, "Peer Marketplace  ·  Donation Centres  ·  Verified Trust", size=14, color=CREAM)
    add_para(tf, "International product presentation  ·  2026", size=12, color=SOFT, space_before=4)
    tb = textbox(s, Inches(10.2), Inches(6.35), Inches(2.5), Inches(0.5))
    set_first(tb.text_frame, "CONFIDENTIAL", size=12, bold=True, color=WARM, align=PP_ALIGN.RIGHT)

    # 2 Agenda
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg_cream(s, W, H); footer(s, 2, W)
    section_label(s, "01  Agenda"); title(s, "What we will cover")
    for i, (n, h, d) in enumerate([
        ("01", "The opportunity", "Why teens need a safer circular fashion platform"),
        ("02", "Product solution", "Marketplace, Donation Hub, and verification"),
        ("03", "Experience design", "Mobile–desktop parity and PWA install"),
        ("04", "Trust & operations", "Admin ERP, centres, and donor motivation"),
        ("05", "Technology & roadmap", "Stack, readiness, and next horizon"),
    ]):
        y = Inches(2.2) + Inches(i * 0.85)
        round_rect(s, Inches(0.7), y, Inches(11.9), Inches(0.72), CARD)
        tb = textbox(s, Inches(0.95), y + Inches(0.12), Inches(0.7), Inches(0.5))
        set_first(tb.text_frame, n, size=20, bold=True, color=FOREST, font="Georgia")
        tb = textbox(s, Inches(1.8), y + Inches(0.08), Inches(9.5), Inches(0.55))
        tf = set_first(tb.text_frame, h, size=18, bold=True, color=FOREST)
        add_para(tf, d, size=13, color=MUTED, space_before=2)

    # 3 Problem
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg_cream(s, W, H); footer(s, 3, W)
    section_label(s, "02  The problem")
    title(s, "Teens outgrow clothes. Trust does not scale.")
    subtitle(s, "Existing marketplaces feel adult-oriented. Donation often means “bag it and hope.”")
    for i, (h, d) in enumerate([
        ("Fast fashion churn", "Clothes are discarded early with no safe local reuse path."),
        ("Trust gap", "Fake photos, unclear sellers, and weak verification erode confidence."),
        ("Donation friction", "Giving clothes to orphanages and centres is opaque and informal."),
        ("Age-fit UX", "Adult marketplaces are not designed for teen clarity or safety."),
    ]):
        x = Inches(0.7) + Inches((i % 2) * 6.2)
        y = Inches(2.35) + Inches((i // 2) * 2.1)
        round_rect(s, x, y, Inches(5.9), Inches(1.85), CARD)
        rect(s, x, y, Inches(0.12), Inches(1.85), FOREST)
        tb = textbox(s, x + Inches(0.4), y + Inches(0.35), Inches(5.2), Inches(1.2))
        tf = set_first(tb.text_frame, h, size=20, bold=True, color=FOREST, font="Georgia")
        add_para(tf, d, size=14, color=MUTED, space_before=8)

    # 4 Solution
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg_cream(s, W, H); footer(s, 4, W)
    section_label(s, "03  Solution")
    title(s, "One platform. Two missions. Visible trust.")
    subtitle(s, "Zyra helps teens reuse clothes locally through a verified peer marketplace and free donation hub.")
    for i, (h, sub, d) in enumerate([
        ("Marketplace", "Sell & buy verified fashion", "Multi-angle photos, COD or online pay, clear fee breakdown."),
        ("Donation Hub", "Give clothes a next home", "Peer gifts + orphanages & centres, multi-item donation basket."),
        ("Verified by Zyra", "Trust you can see", "Admin review before anything goes live — Pending vs Verified."),
    ]):
        x = Inches(0.7) + Inches(i * 4.15)
        round_rect(s, x, Inches(2.4), Inches(3.95), Inches(4.0), FOREST if i == 2 else CARD)
        title_c, sub_c, body_c = (WHITE, SOFT, CREAM) if i == 2 else (FOREST, ACCENT, MUTED)
        tb = textbox(s, x + Inches(0.3), Inches(2.7), Inches(3.35), Inches(3.4))
        tf = set_first(tb.text_frame, f"0{i+1}", size=14, bold=True, color=sub_c)
        add_para(tf, h, size=24, bold=True, color=title_c, space_before=12, font="Georgia")
        add_para(tf, sub, size=14, bold=True, color=sub_c, space_before=10)
        add_para(tf, d, size=14, color=body_c, space_before=14)

    # 5 Marketplace
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg_cream(s, W, H); footer(s, 5, W)
    section_label(s, "04  Marketplace")
    title(s, "A fashion-native buying & selling experience")
    subtitle(s, "International product-card standards: multi-column grids, 3:4 portrait tiles, quick view.")
    for i, (h, d) in enumerate([
        ("Sell composer", "Single-screen listing with required photo angles and pricing."),
        ("Verified feed", "Only Verified by Zyra items appear on Market."),
        ("Buy with clarity", "COD or Stripe test checkout with transparent fees."),
        ("Activity", "Track listings, offers, and transaction status in one place."),
    ]):
        y = Inches(2.25) + Inches(i * 1.05)
        round_rect(s, Inches(0.7), y, Inches(6.3), Inches(0.92), CARD)
        tb = textbox(s, Inches(0.95), y + Inches(0.15), Inches(5.8), Inches(0.7))
        tf = set_first(tb.text_frame, h, size=16, bold=True, color=FOREST)
        add_para(tf, d, size=13, color=MUTED, space_before=2)
    round_rect(s, Inches(7.3), Inches(2.25), Inches(5.3), Inches(4.2), FOREST)
    tb = textbox(s, Inches(7.6), Inches(2.6), Inches(4.7), Inches(3.6))
    tf = set_first(tb.text_frame, "Desktop & mobile parity", size=20, bold=True, color=WHITE, font="Georgia")
    for line in ["Responsive market grids", "Portrait product cards", "Hover / quick-view actions", "Same features on phone & desktop", "Installable PWA shell"]:
        add_para(tf, "→  " + line, size=15, color=CREAM, space_before=12)

    # 6 Donation
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg_cream(s, W, H); footer(s, 6, W)
    section_label(s, "05  Donation Hub")
    title(s, "Give to people — and to places")
    subtitle(s, "Peer gifts plus verified donation centres (orphanages, shelters, community orgs).")
    for i, (h, d) in enumerate([
        ("Centres directory", "Browse verified orphanages and centres with story, needs, and live tagged items."),
        ("Donation basket", "Add multiple clothing items in one session, tag a centre, submit together."),
        ("Claim & fulfill", "Centre staff or peers claim gifts; donors approve in Activity."),
        ("Impact loop", "Donations feed motivation, leaderboard, and public recognition."),
    ]):
        x = Inches(0.7) + Inches((i % 2) * 6.2)
        y = Inches(2.3) + Inches((i // 2) * 2.15)
        round_rect(s, x, y, Inches(5.9), Inches(1.95), CARD)
        tb = textbox(s, x + Inches(0.35), y + Inches(0.35), Inches(5.2), Inches(1.4))
        tf = set_first(tb.text_frame, h, size=20, bold=True, color=FOREST, font="Georgia")
        add_para(tf, d, size=14, color=MUTED, space_before=10)

    # 7 Motivation
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg_cream(s, W, H); footer(s, 7, W)
    section_label(s, "06  Motivation & identity")
    title(s, "Donor tiers and expressive profiles")
    subtitle(s, "Recognition that makes giving stick — without turning the product into a badge farm.")
    for i, (h, d) in enumerate([("Seedling", "Start giving"), ("Helper", "Build rhythm"), ("Guardian", "Consistent impact"), ("Champion", "Community leader")]):
        x = Inches(0.7) + Inches(i * 3.1)
        round_rect(s, x, Inches(2.3), Inches(2.95), Inches(2.2), FOREST if i == 3 else CARD)
        tc, dc = (WHITE, CREAM) if i == 3 else (FOREST, MUTED)
        tb = textbox(s, x + Inches(0.2), Inches(2.7), Inches(2.55), Inches(1.5))
        tf = set_first(tb.text_frame, h, size=18, bold=True, color=tc, font="Georgia", align=PP_ALIGN.CENTER)
        add_para(tf, d, size=13, color=dc, space_before=10, align=PP_ALIGN.CENTER)
    round_rect(s, Inches(0.7), Inches(4.8), Inches(11.9), Inches(1.7), CARD)
    tb = textbox(s, Inches(1.0), Inches(5.05), Inches(11.3), Inches(1.3))
    tf = set_first(tb.text_frame, "Tumblr-like customizable profiles", size=18, bold=True, color=FOREST, font="Georgia")
    add_para(tf, "Banner, accent colour, bio, and layout themes — teens express identity while their public giving story stays visible.", size=14, color=MUTED, space_before=8)

    # 8 Trust
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg_cream(s, W, H); footer(s, 8, W)
    section_label(s, "07  Trust system")
    title(s, "Verification is the product, not a footnote")
    for i, (n, h, d) in enumerate([
        ("01", "List", "Seller / donor submits item with required photos"),
        ("02", "Pending", "Listing waits in admin queue — not public yet"),
        ("03", "Verify", "Zyra admin approves or rejects with reason"),
        ("04", "Live", "Only Verified by Zyra appears on Market / Donate"),
    ]):
        x = Inches(0.7) + Inches(i * 3.1)
        hi = i in (1, 2)
        round_rect(s, x, Inches(2.35), Inches(2.95), Inches(3.9), FOREST if hi else CARD)
        tc, dc, ac = (WHITE, CREAM, WARM) if hi else (FOREST, MUTED, ACCENT)
        tb = textbox(s, x + Inches(0.25), Inches(2.7), Inches(2.45), Inches(3.2))
        tf = set_first(tb.text_frame, n, size=14, bold=True, color=ac)
        add_para(tf, h, size=22, bold=True, color=tc, space_before=14, font="Georgia")
        add_para(tf, d, size=14, color=dc, space_before=16)

    # 9 Journeys
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg_cream(s, W, H); footer(s, 9, W)
    section_label(s, "08  User journeys"); title(s, "Three primary paths")
    for i, (h, steps) in enumerate([
        ("Seller", ["Open Sell composer", "Add photos + price", "Submit → Pending", "Verified → live on Market", "Buyer pays COD / online"]),
        ("Donor", ["Browse centres", "Build donation basket", "Tag orphanage / centre", "Zyra verifies gifts", "Staff or peer claims"]),
        ("Admin", ["Open ERP queue", "Review photos & details", "Approve or reject", "Track transactions", "Manage centres & settings"]),
    ]):
        x = Inches(0.7) + Inches(i * 4.15)
        round_rect(s, x, Inches(2.25), Inches(3.95), Inches(4.25), CARD)
        rect(s, x, Inches(2.25), Inches(3.95), Inches(0.7), FOREST)
        tb = textbox(s, x + Inches(0.2), Inches(2.38), Inches(3.55), Inches(0.5))
        set_first(tb.text_frame, h, size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER, font="Georgia")
        tb = textbox(s, x + Inches(0.3), Inches(3.15), Inches(3.35), Inches(3.1))
        tf = set_first(tb.text_frame, steps[0], size=14, color=FOREST)
        for step in steps[1:]:
            add_para(tf, step, size=14, color=MUTED, space_before=10)

    # 10 Tech
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg_cream(s, W, H); footer(s, 10, W)
    section_label(s, "09  Technology")
    title(s, "Production-grade stack, live data")
    subtitle(s, "Not a disposable mock — every listing, claim, and verification hits live Supabase.")
    for i, (h, d) in enumerate([
        ("Frontend", "Next.js App Router, TypeScript, Tailwind, shadcn/ui"),
        ("Backend data", "Supabase Postgres + Auth + Storage + RLS"),
        ("Payments", "Stripe test mode + COD-ready checkout flows"),
        ("Delivery", "Responsive web + Progressive Web App install"),
        ("Ops", "Admin ERP: queue, transactions, CMS, analytics, centres"),
        ("Quality", "Mobile–desktop feature parity as a product rule"),
    ]):
        x = Inches(0.7) + Inches((i % 3) * 4.15)
        y = Inches(2.35) + Inches((i // 3) * 2.15)
        round_rect(s, x, y, Inches(3.95), Inches(1.95), CARD)
        tb = textbox(s, x + Inches(0.3), y + Inches(0.35), Inches(3.35), Inches(1.4))
        tf = set_first(tb.text_frame, h, size=18, bold=True, color=FOREST, font="Georgia")
        add_para(tf, d, size=14, color=MUTED, space_before=10)

    # 11 Admin
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg_cream(s, W, H); footer(s, 11, W)
    section_label(s, "10  Operations"); title(s, "Admin ERP to run the live platform")
    for i, (h, d) in enumerate([
        ("Verification queue", "Approve / reject listings with audit reasons"),
        ("Transactions", "Monitor buys, COD, and seller payout actions"),
        ("Centres admin", "Manage donation destinations and visibility"),
        ("CMS & experiments", "Heroes, content blocks, A/B tests"),
        ("Analytics", "Audience, listing, and seller performance views"),
        ("Settings", "Currency, fee %, users, and approvers"),
    ]):
        x = Inches(0.7) + Inches((i % 3) * 4.15)
        y = Inches(2.3) + Inches((i // 3) * 2.15)
        hi = (i // 3) == 0
        round_rect(s, x, y, Inches(3.95), Inches(1.95), FOREST if hi else CARD)
        tc, dc = (WHITE, CREAM) if hi else (FOREST, MUTED)
        tb = textbox(s, x + Inches(0.3), y + Inches(0.4), Inches(3.35), Inches(1.3))
        tf = set_first(tb.text_frame, h, size=17, bold=True, color=tc, font="Georgia")
        add_para(tf, d, size=13, color=dc, space_before=10)

    # 12 Positioning
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg_cream(s, W, H); footer(s, 12, W)
    section_label(s, "11  Positioning"); title(s, "Why Zyra stands apart")
    rect(s, Inches(0.7), Inches(2.2), Inches(11.9), Inches(0.55), FOREST)
    for i, h in enumerate(["Dimension", "Zyra", "Typical alternatives"]):
        x = Inches(0.9) if i == 0 else (Inches(3.5) if i == 1 else Inches(8.2))
        tb = textbox(s, x, Inches(2.3), Inches(4), Inches(0.4))
        set_first(tb.text_frame, h, size=14, bold=True, color=WHITE)
    for i, (a, b, c) in enumerate([
        ("Audience", "Teen-clear language and safety-first flows", "Adult resale marketplaces"),
        ("Trust", "Mandatory admin verification before live", "Self-serve, trust optional"),
        ("Impact", "Marketplace + donation centres in one product", "Sell-only or charity-only apps"),
        ("Identity", "Custom profiles + donor tiers", "Generic account pages"),
        ("Delivery", "PWA + mobile–desktop parity", "Mobile-only or desktop-first gaps"),
    ]):
        y = Inches(2.75) + Inches(i * 0.75)
        rect(s, Inches(0.7), y, Inches(11.9), Inches(0.75), CARD if i % 2 == 0 else WARM)
        for j, txt in enumerate([a, b, c]):
            x = Inches(0.9) if j == 0 else (Inches(3.5) if j == 1 else Inches(8.2))
            w = Inches(2.4) if j == 0 else Inches(4.4)
            tb = textbox(s, x, y + Inches(0.18), w, Inches(0.45))
            set_first(tb.text_frame, txt, size=13, bold=(j == 0), color=FOREST if j < 2 else MUTED)

    # 13 Roadmap
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg_cream(s, W, H); footer(s, 13, W)
    section_label(s, "12  Roadmap"); title(s, "Shipped now. Next horizon.")
    for i, (h, bg, tc, dc, items) in enumerate([
        ("Production-capable today", FOREST, WHITE, CREAM, [
            "Verified marketplace & donation hub", "Donation centres + basket",
            "Donor tiers & themed profiles", "Admin ERP modules", "PWA installable shell"]),
        ("Near-term hardening", CARD, FOREST, MUTED, [
            "E2E test suite & rate limits", "Richer messaging UI pages",
            "Centre self-signup / KYC", "Push on tier milestones"]),
        ("Strategic expansion", CARD, FOREST, MUTED, [
            "Money donations to centres", "Logistics partnerships",
            "Live payments & payouts", "Native app wrappers"]),
    ]):
        x = Inches(0.7) + Inches(i * 4.15)
        round_rect(s, x, Inches(2.25), Inches(3.95), Inches(4.25), bg)
        tb = textbox(s, x + Inches(0.25), Inches(2.55), Inches(3.45), Inches(3.7))
        tf = set_first(tb.text_frame, h, size=16, bold=True, color=tc, font="Georgia")
        for it in items:
            add_para(tf, "•  " + it, size=14, color=dc, space_before=12)

    # 14 Closing
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, FOREST)
    rect(s, Inches(0.7), Inches(2.4), Inches(1.4), Inches(0.08), WARM)
    tb = textbox(s, Inches(0.7), Inches(1.5), Inches(12), Inches(0.5))
    set_first(tb.text_frame, "CLOSING", size=14, bold=True, color=WARM)
    tb = textbox(s, Inches(0.7), Inches(2.7), Inches(11.5), Inches(1.5))
    tf = set_first(tb.text_frame, "Clothes that get a second life.", size=36, bold=True, color=WHITE, font="Georgia")
    add_para(tf, "Zyra turns outgrown fashion into trusted local reuse — for teens, centres, and communities.", size=16, color=CREAM, space_before=16)
    tb = textbox(s, Inches(0.7), Inches(5.0), Inches(11), Inches(1.5))
    tf = set_first(tb.text_frame, "Next step", size=14, bold=True, color=WARM)
    add_para(tf, "Product walkthrough  ·  Demo accounts  ·  Pilot with a donation centre", size=16, color=WHITE, space_before=8)
    add_para(tf, "zyra  ·  circular fashion platform", size=12, color=SOFT, space_before=18)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"Wrote {OUT} ({OUT.stat().st_size} bytes, {len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
